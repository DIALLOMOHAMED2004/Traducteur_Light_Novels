import os
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape
from docx import Document
from PIL import Image
from PyPDF2 import PdfWriter

from .forms import DocumentForm, MAX_UPLOAD_SIZE
from .models import DocFile, Subscriber, TranslationJob
from .services.workspace import workspace_for
from .usage_policy import FREE_POLICY, PAID_POLICY, get_usage_policy


def pdf_upload(pages=1, name='source.pdf'):
    buffer = BytesIO()
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    writer.write(buffer)
    return SimpleUploadedFile(name, buffer.getvalue(), content_type='application/pdf')


def image_upload(name='source.png', size=(8, 8)):
    buffer = BytesIO()
    Image.new('RGB', size, color='white').save(
        buffer, format='JPEG' if name.endswith('.jpg') else 'PNG',
    )
    return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/png')


class UsagePolicyTests(SimpleTestCase):
    def test_current_plan_selects_unchanged_limits(self):
        user = SimpleNamespace(state_abonnement=False)
        for paid, expected, limits in [(False, FREE_POLICY, (2, 5, 3)),
                                       (True, PAID_POLICY, (10, 30, 15)),
                                       (False, FREE_POLICY, (2, 5, 3))]:
            with self.subTest(paid=paid):
                user.state_abonnement = paid
                policy = get_usage_policy(user)
                self.assertIs(policy, expected)
                self.assertEqual((policy.max_pdf, policy.max_pdf_pages, policy.max_images), limits)


class DocumentFormTests(SimpleTestCase):
    def form(self, upload, kind='PDF', language='Anglais'):
        return DocumentForm(
            {'type_file': kind, 'type_language': language},
            {'button_televerse': upload},
        )

    def test_valid_pdf_jpg_png_and_stream_rewound(self):
        for upload, kind in [(pdf_upload(), 'PDF'), (pdf_upload(name='source.PDF'), 'PDF'),
                             (image_upload(), 'Image'), (image_upload('source.jpg'), 'Image')]:
            with self.subTest(name=upload.name):
                form = self.form(upload, kind)
                self.assertTrue(form.is_valid(), form.errors)
                self.assertEqual(form.cleaned_data['button_televerse'].tell(), 0)

    def test_invalid_extension(self):
        for upload, kind in [(pdf_upload(name='source.txt'), 'PDF'),
                             (image_upload('source.gif'), 'Image'),
                             (image_upload('source.JPG'), 'Image')]:
            with self.subTest(name=upload.name):
                self.assertIn('button_televerse', self.form(upload, kind).errors)

    def test_declared_type_and_content_must_match(self):
        for upload, kind in [(image_upload(), 'PDF'), (pdf_upload(), 'Image'),
                             (image_upload('source.pdf'), 'PDF'),
                             (pdf_upload(name='source.png'), 'Image')]:
            with self.subTest(name=upload.name, kind=kind):
                self.assertIn('button_televerse', self.form(upload, kind).errors)

    def test_corrupt_pdf_and_image(self):
        for name, kind in [('bad.pdf', 'PDF'), ('bad.png', 'Image')]:
            with self.subTest(kind=kind):
                form = self.form(SimpleUploadedFile(name, b'not a document'), kind)
                self.assertIn('button_televerse', form.errors)

    def test_pdf_without_pages(self):
        self.assertIn('aucune page', str(self.form(pdf_upload(0)).errors))

    def test_size_limit_inclusive(self):
        for kind in ('PDF', 'Image'):
            for size, allowed in [(MAX_UPLOAD_SIZE, True), (MAX_UPLOAD_SIZE + 1, False)]:
                with self.subTest(kind=kind, size=size):
                    upload = pdf_upload() if kind == 'PDF' else image_upload()
                    upload.size = size
                    form = self.form(upload, kind)
                    self.assertEqual(form.is_valid(), allowed, form.errors)

    def test_image_dimension_and_pixel_boundaries(self):
        # Simuler l'en-tête évite d'allouer des images de dizaines de millions de pixels.
        for size, allowed in [((20000, 1), True), ((20001, 1), False),
                              ((1, 20001), False), ((8000, 5000), True),
                              ((8000, 5001), False)]:
            with self.subTest(size=size), patch('Tsukiyomi_app.forms.Image.open') as opened:
                opened.return_value.__enter__.return_value.size = size
                form = self.form(SimpleUploadedFile('image.png', b'header'), 'Image')
                self.assertEqual(form.is_valid(), allowed, form.errors)

    def test_image_decompression_bomb_is_rejected(self):
        with patch('Tsukiyomi_app.forms.Image.open', side_effect=Image.DecompressionBombError):
            self.assertIn('button_televerse', self.form(image_upload(), 'Image').errors)

    def test_language_and_type_choices_are_validated(self):
        self.assertIn('type_language', self.form(pdf_upload(), language='unknown').errors)
        self.assertIn('type_file', self.form(pdf_upload(), kind='unknown').errors)


class UploadTests(TestCase):
    """Tester les réponses HTTP et artefacts dans un espace entièrement temporaire."""

    def setUp(self):
        temp = TemporaryDirectory(prefix='tsukiyomi-test-')
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        previous = Path.cwd()
        os.chdir(self.root)
        self.addCleanup(os.chdir, previous)
        config = self.settings(MEDIA_ROOT=self.root / 'media_upload',
                               EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
        config.enable()
        self.addCleanup(config.disable)
        self.user = get_user_model().objects.create_user(
            username='reader', email='reader@example.test', password='test-password',
        )
        self.client.force_login(self.user)
        self.start_patch('socket.socket.connect', side_effect=AssertionError('Réseau interdit dans ces tests'))
        self.start_patch('subprocess.Popen', side_effect=AssertionError('Processus externe interdit dans ces tests'))
        self.convert = self.start_patch('Tsukiyomi_app.services.extraction.convert_from_path', return_value=['page'])
        self.ocr = self.start_patch('Tsukiyomi_app.services.ocr.pytesseract.image_to_string', return_value='Texte source')
        self.translator = self.start_patch('Tsukiyomi_app.services.translation.GoogleTranslator')
        self.translator.return_value.translate.return_value = 'Texte traduit'
        self.send = self.start_patch('Tsukiyomi_app.services.notification.EmailMessage.send', autospec=True, return_value=1)
        self.mass_mail = self.start_patch('Tsukiyomi_app.views.send_mass_mail', return_value=1)

    def start_patch(self, target, **kwargs):
        patcher = patch(target, **kwargs)
        result = patcher.start()
        self.addCleanup(patcher.stop)
        return result

    def use_plan(self, paid):
        self.user.state_abonnement = paid
        self.user.save(update_fields=['state_abonnement'])
        return reverse('televerse_url')

    def upload(self, paid=False, kind='PDF', pages=1, language='Anglais', file=None):
        self.convert.return_value = [f'page-{i}' for i in range(pages)]
        return self.client.post(self.use_plan(paid), {
            'profileType': 'pdf_img', 'pi-type_file': kind,
            'pi-type_language': language,
            'pi-button_televerse': file or (pdf_upload(pages) if kind == 'PDF' else image_upload()),
        })

    def existing_documents(self, count, kind='PDF', user=None):
        return [DocFile.objects.create(
            user=user or self.user, type_file=kind, type_language='Anglais',
            button_televerse=f'media/old-{kind}-{i}',
        ) for i in range(count)]

    def test_anonymous_cannot_start_pipeline(self):
        self.client.logout()
        for route in ['reindex', 'televerse_url', 'name_televerse_url_free', 'name_televerse_url_paid']:
            for method in ['get', 'post']:
                with self.subTest(route=route, method=method):
                    response = getattr(self.client, method)(reverse(route), {
                        'profileType': 'pdf_img', 'pi-type_file': 'PDF',
                        'pi-type_language': 'Anglais', 'pi-button_televerse': pdf_upload(),
                    })
                    self.assertEqual(response.status_code, 302)
                    self.assertTrue(response.url.startswith(reverse('login')))
        self.assertEqual(DocFile.objects.count(), 0)
        for dependency in [self.convert, self.ocr, self.translator, self.send]:
            dependency.assert_not_called()
        self.assertFalse((self.root / 'Tsukiyomi_doc').exists())

    def test_common_form_and_legacy_access_restrictions(self):
        for paid in [False, True]:
            with self.subTest(paid=paid):
                target = self.use_plan(paid)
                legacy = reverse('name_televerse_url_paid' if paid else 'name_televerse_url_free')
                for url in [target, legacy]:
                    response = self.client.get(url)
                    self.assertContains(response, '<form ', count=1)
                    self.assertContains(response, f'action="{target}"')
                    self.assertContains(response, 'name="csrfmiddlewaretoken"')
                    self.assertContains(response, 'name="pi-button_televerse"')
                    self.assertTemplateUsed(response, 'tsukiyomi_app/televerse_page.html')
                wrong = reverse('name_televerse_url_free' if paid else 'name_televerse_url_paid')
                for method in ['get', 'post']:
                    response = getattr(self.client, method)(wrong, {
                        'profileType': 'pdf_img', 'pi-type_file': 'PDF',
                        'pi-type_language': 'Anglais', 'pi-button_televerse': pdf_upload(),
                    })
                    self.assertEqual(response.status_code, 200)
                    self.assertTemplateUsed(response, 'tsukiyomi_app/error_abonnement_free.html' if paid
                                            else 'tsukiyomi_app/error_abonnement_payant.html')
        for dependency in [self.convert, self.ocr, self.translator, self.send]:
            dependency.assert_not_called()
        self.assertEqual(DocFile.objects.count(), 0)

    def test_legacy_posts_share_pipeline_and_usage_with_common_route(self):
        for paid, kind, limit in [(False, 'PDF', 2), (True, 'PDF', 10),
                                  (False, 'Image', 3), (True, 'Image', 15)]:
            with self.subTest(paid=paid, kind=kind):
                DocFile.objects.all().delete()
                self.use_plan(paid)
                self.existing_documents(limit - 1, kind)
                legacy = reverse('name_televerse_url_paid' if paid else 'name_televerse_url_free')
                self.send.reset_mock()
                response = self.client.post(legacy, {
                    'profileType': 'pdf_img', 'pi-type_file': kind,
                    'pi-type_language': 'Anglais',
                    'pi-button_televerse': pdf_upload() if kind == 'PDF' else image_upload(),
                })
                self.assertContains(response, 'Texte traduit')
                self.assertEqual(DocFile.objects.count(), limit)
                self.send.assert_called_once()
                doc = Document(BytesIO(self.send.call_args.args[0].attachments[0][1]))
                self.assertIn('Texte traduit', [p.text for p in doc.paragraphs])
                self.ocr.reset_mock()
                self.send.reset_mock()
                response = self.upload(paid, kind)
                self.assertContains(response, 'Limite de téléversement')
                self.assertEqual(DocFile.objects.count(), limit)
                self.ocr.assert_not_called()
                self.send.assert_not_called()

    def test_pdf_and_image_quota_boundaries(self):
        for paid, kind, limit in [(False, 'PDF', 2), (True, 'PDF', 10),
                                  (False, 'Image', 3), (True, 'Image', 15)]:
            with self.subTest(paid=paid, kind=kind):
                DocFile.objects.all().delete()
                self.existing_documents(limit - 1, kind)
                response = self.upload(paid, kind)
                self.assertEqual(response.status_code, 200)
                self.assertTemplateUsed(response, 'tsukiyomi_app/succes_uploadfile.html')
                self.assertEqual(DocFile.objects.count(), limit)
                self.ocr.reset_mock()
                self.send.reset_mock()
                response = self.upload(paid, kind)
                self.assertContains(response, 'Limite de téléversement')
                self.assertEqual(DocFile.objects.count(), limit)
                self.ocr.assert_not_called()
                self.send.assert_not_called()

    def test_pdf_page_boundaries(self):
        for paid, limit in [(False, 5), (True, 30)]:
            with self.subTest(paid=paid):
                DocFile.objects.all().delete()
                self.assertEqual(self.upload(paid, pages=limit).status_code, 200)
                self.assertEqual(DocFile.objects.count(), 1)
                self.ocr.reset_mock()
                response = self.upload(paid, pages=limit + 1)
                self.assertContains(response, 'pages')
                self.assertTemplateUsed(response, 'tsukiyomi_app/televerse_page.html')
                self.assertEqual(DocFile.objects.count(), 1)
                self.ocr.assert_not_called()

    def test_quotas_are_per_user_and_file_type(self):
        other = get_user_model().objects.create_user(username='other')
        self.existing_documents(20, user=other)
        self.existing_documents(3, kind='Image')
        response = self.upload()
        self.assertTemplateUsed(response, 'tsukiyomi_app/succes_uploadfile.html')
        self.assertEqual(DocFile.objects.filter(user=self.user, type_file='PDF').count(), 1)

    def test_invalid_upload_does_not_save_or_start_pipeline(self):
        for paid in [False, True]:
            with self.subTest(paid=paid):
                response = self.upload(paid, file=SimpleUploadedFile('bad.pdf', b'bad'))
                self.assertTrue(response.context['DocForm'].errors)
        self.assertEqual(DocFile.objects.count(), 0)
        for dependency in [self.convert, self.ocr, self.translator, self.send]:
            dependency.assert_not_called()

    def test_empty_ocr_stops_before_translation_and_email(self):
        for paid in [False, True]:
            for kind in ['PDF', 'Image']:
                for text in ['', None]:
                    with self.subTest(paid=paid, kind=kind, text=text):
                        DocFile.objects.all().delete()
                        self.ocr.return_value = text
                        response = self.upload(paid, kind)
                        self.assertContains(response, 'aucun text')
                        self.translator.assert_not_called()
                        self.send.assert_not_called()

    def test_media_deletion_is_restricted_without_touching_real_files(self):
        sentinel = self.root / 'media_upload/media/sentinel.pdf'
        sentinel.parent.mkdir(parents=True)
        sentinel.write_bytes(b'test-only')
        for anonymous in [False, True]:
            if anonymous:
                self.client.logout()
            response = self.client.post(reverse('clear_media_name'))
            self.assertEqual(response.status_code, 302)
            self.assertTrue(sentinel.exists())
        self.user.is_superuser = True
        self.user.save(update_fields=['is_superuser'])
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('clear_media_name')).status_code, 405)
        self.assertTrue(sentinel.exists())


    def test_newly_uploaded_document_is_processed_among_existing_documents(self):
        for kind in ['PDF', 'Image']:
            with self.subTest(kind=kind):
                DocFile.objects.all().delete()
                timestamp = timezone.now()
                old = [DocFile.objects.create(
                    user=self.user, type_file=kind, type_language='Anglais',
                    button_televerse=pdf_upload(name=f'old-{i}.pdf') if kind == 'PDF'
                    else image_upload(f'old-{i}.png'),
                ) for i in range(2)]
                DocFile.objects.filter(pk__in=[doc.pk for doc in old]).update(uploaded_at=timestamp)
                # Des dates identiques et un ordre légal différent ne doivent pas changer la source.
                ordering = ['-pk'] if kind == 'PDF' else ['pk']
                with patch.object(DocFile._meta, 'ordering', ordering), patch(
                    'django.utils.timezone.now', return_value=timestamp,
                ):
                    response = self.upload(paid=True, kind=kind)
                self.assertEqual(response.status_code, 200)
                current = DocFile.objects.exclude(pk__in=[doc.pk for doc in old]).get()
                actual_path = (self.convert.call_args.args[0] if kind == 'PDF'
                               else self.ocr.call_args.args[0].filename)
                job = TranslationJob.objects.get(source=current)
                self.assertEqual(Path(actual_path).parent, workspace_for(job).source_dir)
                self.assertEqual(Path(actual_path).read_bytes(), Path(current.button_televerse.path).read_bytes())
                self.assertEqual(self.send.call_args.args[0].to, [self.user.email])

    def test_docx_pages_remain_in_numeric_order_and_exclude_previous_results(self):
        for paid, count in [(False, 3), (True, 12)]:
            with self.subTest(paid=paid):
                DocFile.objects.all().delete()
                folder = self.root / f'Tsukiyomi_doc/repository-{self.user}'
                folder.mkdir(parents=True, exist_ok=True)
                # Un ancien résultat et une page résiduelle ne font pas partie du nouveau PDF.
                for name in ['trad_fusion_reader.docx', 'trad_fr99.docx']:
                    old = Document()
                    old.add_paragraph('ANCIEN DOCUMENT')
                    old.save(folder / name)
                texts = [f'PAGE {i}' for i in range(1, count + 1)]
                self.ocr.side_effect = texts
                self.translator.return_value.translate.side_effect = lambda text, **kwargs: text
                listdir = os.listdir

                def unordered_listdir(path):
                    names = listdir(path)
                    return sorted(names, reverse=True) if Path(path) == Path(folder.relative_to(self.root)) else names

                with patch('os.listdir', side_effect=unordered_listdir):
                    response = self.upload(paid, pages=count)
                self.assertEqual(response.status_code, 200)
                email = self.send.call_args.args[0]
                self.assertEqual(email.to, [self.user.email])
                self.assertEqual(len(email.attachments), 1)
                attachment = email.attachments[0]
                doc = Document(BytesIO(attachment[1]))
                paragraphs = [paragraph.text for paragraph in doc.paragraphs
                              if paragraph.text != 'TRADUIT PAR ZENIA']
                self.assertEqual(paragraphs, texts)
                self.assertTrue(attachment[0].endswith('.docx'))
                self.assertIn('Français', email.subject)
                self.assertFalse(self.send.call_args.kwargs['fail_silently'])

    def test_image_docx_contains_translation_and_is_attached(self):
        for paid in [False, True]:
            for name in ['source.png', 'source.jpg']:
                with self.subTest(paid=paid, name=name):
                    DocFile.objects.all().delete()
                    response = self.upload(paid, kind='Image', file=image_upload(name))
                    self.assertContains(response, 'Texte traduit')
                    email = self.send.call_args.args[0]
                    self.assertEqual(email.to, [self.user.email])
                    self.assertEqual(email.from_email, settings.DEFAULT_FROM_EMAIL)
                    doc = Document(BytesIO(email.attachments[0][1]))
                    self.assertEqual([p.text for p in doc.paragraphs], ['TRADUIT PAR ZENIA', 'Texte traduit'])



    def test_all_six_languages_work_for_pdf_and_image_in_both_plans(self):
        from deep_translator import GoogleTranslator
        self.translator.side_effect = GoogleTranslator
        for paid in [False, True]:
            for kind in ['PDF', 'Image']:
                for language, ocr_language in [('Français', 'fra'), ('Anglais', 'eng'),
                                               ('Italien', 'ita'), ('Espagnol', 'spa'),
                                               ('Japonais', 'jpn'), ('Chinois', 'chi_tra')]:
                    with self.subTest(paid=paid, kind=kind, language=language):
                        DocFile.objects.all().delete()
                        self.send.reset_mock()
                        # Le constructeur réel valide la langue sans réseau ; seule la traduction est simulée.
                        with patch.object(GoogleTranslator, 'translate', return_value='Texte traduit'):
                            response = self.upload(paid, kind, language=language)
                        self.assertContains(response, 'Texte traduit')
                        self.assertEqual(self.ocr.call_args.kwargs['lang'], ocr_language)
                        self.send.assert_called_once()
                        doc = Document(BytesIO(self.send.call_args.args[0].attachments[0][1]))
                        self.assertIn('Texte traduit', [p.text for p in doc.paragraphs])

    def test_ocr_errors_stop_processing_and_are_visible(self):
        from pytesseract.pytesseract import TesseractError, TesseractNotFoundError
        for paid in [False, True]:
            for kind in ['PDF', 'Image']:
                for error in [TesseractError(1, 'failure'), TesseractNotFoundError()]:
                    with self.subTest(paid=paid, kind=kind, error=type(error).__name__):
                        DocFile.objects.all().delete()
                        self.ocr.side_effect = error
                        response = self.upload(paid, kind)
                        self.assertContains(response, 'traitement', status_code=502)
                        self.assertTemplateUsed(response, 'tsukiyomi_app/televerse_page.html')
                        self.translator.assert_not_called()
                        self.send.assert_not_called()
                        self.assertFalse(list(self.root.rglob('*.docx')))
                        self.assertEqual(DocFile.objects.count(), 1)

    def test_translation_errors_stop_before_rendering_and_email(self):
        from deep_translator.exceptions import RequestError, TooManyRequests, TranslationNotFound
        from requests.exceptions import ConnectionError as ProviderConnectionError, Timeout
        errors = [ConnectionError('offline'), ProviderConnectionError('offline'), Timeout('timeout'),
                  RequestError(), TooManyRequests(), TranslationNotFound('source')]
        for paid in [False, True]:
            for kind in ['PDF', 'Image']:
                for language in ['Français', 'Anglais']:
                    for error in errors:
                        with self.subTest(paid=paid, kind=kind, language=language, error=type(error).__name__):
                            DocFile.objects.all().delete()
                            self.translator.return_value.translate.side_effect = error
                            response = self.upload(paid, kind, language=language)
                            self.assertContains(response, 'traitement', status_code=502)
                            self.send.assert_not_called()
                            self.assertFalse(list(self.root.rglob('*.docx')))
                            self.assertEqual(DocFile.objects.count(), 1)

    def test_error_on_later_pdf_page_never_sends_partial_result(self):
        from deep_translator.exceptions import RequestError
        for paid in [False, True]:
            with self.subTest(paid=paid):
                DocFile.objects.all().delete()
                self.translator.return_value.translate.side_effect = ['PAGE 1', RequestError()]
                response = self.upload(paid, pages=3)
                self.assertEqual(response.status_code, 502)
                self.send.assert_not_called()
                self.assertFalse(list(self.root.rglob('output/*.docx')))

    def test_poppler_errors_stop_before_ocr(self):
        from pdf2image.exceptions import (PDFInfoNotInstalledError, PDFPageCountError,
                                         PDFPopplerTimeoutError, PDFSyntaxError)
        for paid in [False, True]:
            for error_type in [PDFInfoNotInstalledError, PDFPageCountError, PDFPopplerTimeoutError, PDFSyntaxError]:
                with self.subTest(paid=paid, error=error_type.__name__):
                    DocFile.objects.all().delete()
                    self.convert.side_effect = error_type('failure')
                    response = self.upload(paid)
                    self.assertContains(response, 'traitement', status_code=502)
                    self.ocr.assert_not_called()
                    self.send.assert_not_called()

    def test_email_errors_are_visible_after_docx_generation(self):
        from smtplib import SMTPException
        for paid in [False, True]:
            for kind in ['PDF', 'Image']:
                for error in [SMTPException('rejected'), ConnectionRefusedError('offline')]:
                    with self.subTest(paid=paid, kind=kind, error=type(error).__name__):
                        DocFile.objects.all().delete()
                        self.send.side_effect = error
                        response = self.upload(paid, kind)
                        self.assertContains(response, 'traitement', status_code=502)
                        self.assertTemplateUsed(response, 'tsukiyomi_app/televerse_page.html')
                        email = self.send.call_args.args[0]
                        self.assertEqual(email.to, [self.user.email])
                        self.assertEqual(len(email.attachments), 1)
                        self.assertEqual(DocFile.objects.count(), 1)



    def test_unexpected_programming_error_is_not_hidden(self):
        self.translator.return_value.translate.side_effect = ValueError('unexpected bug')
        with self.assertRaisesRegex(ValueError, 'unexpected bug'):
            self.upload()
        self.send.assert_not_called()

    def test_changing_plan_preserves_existing_usage(self):
        self.existing_documents(2)
        response = self.upload(paid=True)
        self.assertTemplateUsed(response, 'tsukiyomi_app/succes_uploadfile.html')
        self.assertEqual(DocFile.objects.count(), 3)
        self.ocr.reset_mock()
        response = self.upload(paid=False)
        self.assertContains(response, 'Limite de téléversement')
        self.assertEqual(DocFile.objects.count(), 3)
        self.ocr.assert_not_called()


class NewsletterTests(TestCase):
    def setUp(self):
        patcher = patch('Tsukiyomi_app.views.send_mass_mail', return_value=1)
        self.send = patcher.start()
        self.addCleanup(patcher.stop)

    def test_public_subscription_sends_to_one_complete_address(self):
        response = self.client.post(reverse('subscrib_newsletter'), {'email': 'reader@example.test'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Subscriber.objects.get().email, 'reader@example.test')
        self.send.assert_called_once()
        messages = self.send.call_args.args[0]
        self.assertEqual(len(messages), 1)
        self.assertEqual(messages[0][2], settings.DEFAULT_FROM_EMAIL)
        self.assertEqual(messages[0][3], ['reader@example.test'])
        self.assertFalse(self.send.call_args.kwargs['fail_silently'])
        self.assertTrue(response.context['msg_success_newsletter'])

    def test_invalid_or_duplicate_email_does_not_send(self):
        Subscriber.objects.create(email='existing@example.test')
        for address in ['invalid', 'existing@example.test']:
            with self.subTest(address=address):
                response = self.client.post(reverse('subscrib_newsletter'), {'email': address})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['sf'].errors)
        self.assertEqual(Subscriber.objects.count(), 1)
        self.send.assert_not_called()

    def test_email_error_is_visible_and_subscription_is_preserved(self):
        from smtplib import SMTPException
        for error in [ConnectionRefusedError('offline'), SMTPException('rejected')]:
            with self.subTest(error=type(error).__name__):
                Subscriber.objects.all().delete()
                self.send.side_effect = error
                response = self.client.post(reverse('subscrib_newsletter'), {'email': 'reader@example.test'})
                self.assertContains(response, escape("L'inscription est enregistrée"), status_code=502)
                self.assertNotIn('msg_success_newsletter', response.context)
                self.assertEqual(Subscriber.objects.count(), 1)
