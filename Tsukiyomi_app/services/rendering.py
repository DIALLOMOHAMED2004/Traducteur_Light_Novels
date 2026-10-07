from docx import Document


def write_docx(text, path):
    """Créer une page traduite avec le titre historique."""
    document = Document()
    document.add_heading('TRADUIT PAR ZENIA')
    document.add_paragraph(text)
    document.save(path)


def merge_docx(page_paths, output_path):
    """Fusionner uniquement les pages fournies, dans leur ordre explicite."""
    merged = Document()
    for path in page_paths:
        page = Document(path)
        for element in page.element.body:
            merged.element.body.append(element)
    merged.save(output_path)
