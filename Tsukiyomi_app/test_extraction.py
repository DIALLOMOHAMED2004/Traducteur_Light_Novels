from io import BytesIO
from unittest.mock import call, patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from docx import Document
from PIL import Image, ImageDraw
from PyPDF2 import PdfReader, PdfWriter
from PyPDF2.errors import PdfReadError
from PyPDF2.generic import DecodedStreamObject, DictionaryObject, NameObject
from pytesseract.pytesseract import TesseractError

from .models import DocFile, TranslationJob
from .services.errors import DocumentProcessingError
from .services.extraction import usable_native_text
from .services.processor import process_document
from .services.workspace import workspace_for
from .test_jobs import JobTestSetup


def native_pdf_upload(texts):
    """Construire un vrai PDF latin avec texte incorporé ; None crée une page sans texte."""
    writer = PdfWriter()
    for text in texts:
        writer.add_blank_page(width=300, height=300)
        page = writer.pages[-1]
        if text is None:
            continue
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({
                NameObject('/F1'): DictionaryObject({
                    NameObject('/Type'): NameObject('/Font'),
                    NameObject('/Subtype'): NameObject('/Type1'),
                    NameObject('/BaseFont'): NameObject('/Helvetica'),
                    NameObject('/Encoding'): NameObject('/WinAnsiEncoding'),
                }),
            }),
        })
        stream = DecodedStreamObject()
        encoded = text.encode('cp1252').hex()
        stream.set_data(f'BT /F1 12 Tf 10 250 Td <{encoded}> Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = stream
    buffer = BytesIO()
    writer.write(buffer)
    return SimpleUploadedFile('source.pdf', buffer.getvalue(), content_type='application/pdf')


class NativeTextQualityTests(SimpleTestCase):
    def test_readable_text_in_all_six_languages_and_short_pages(self):
        for text in [
            'Épilogue : où étions-nous ?', 'The story begins here.',
            'Capitolo primo: perché?', '¿Dónde está el corazón?',
            '第一章　物語が始まる。', '第一章，故事開始了。',
            'I', '7', '一', '序', 'Prologue', 'Chapitre 1',
            'E\u0301pilogue', '\tDeux lignes.\r\nUne suite.\n',
            'Une\u00a0espace et une cé\u00adsure.',
        ]:
            with self.subTest(text=text):
                self.assertTrue(usable_native_text(text))

    def test_empty_or_obviously_corrupt_text_is_rejected(self):
        for text in [
            None, '', ' \n\r\t', '\u200b', '\u0301\u0301', '??? !!!', 'a %%%% !!!',
            'Texte \ufffd abîmé', 'Texte\x00illisible', '\x01\x02',
            'Texte \ue000 privé', '\ud800', '\u0378',
            '(cid:123) (cid:456)', 'aaaaaaaa', '11111111',
        ]:
            with self.subTest(text=text):
                self.assertFalse(usable_native_text(text))


class AdaptivePDFTests(JobTestSetup, TestCase):
    """Extraction PyPDF2 et DOCX réels ; OCR, Poppler, traduction et SMTP simulés."""

    def native_job(self, texts, language='Anglais'):
        source = DocFile.objects.create(
            user=self.user, type_file='PDF', type_language=language,
            button_televerse=native_pdf_upload(texts),
        )
        return TranslationJob.objects.create(user=self.user, source=source, source_language=language)

    def assert_document(self, job, result, texts):
        job.refresh_from_db()
        self.assertEqual(job.status, TranslationJob.Status.SUCCEEDED)
        self.assertEqual(job.error, '')
        self.assertLessEqual(job.created_at, job.started_at)
        self.assertLessEqual(job.started_at, job.finished_at)
        self.assertEqual(result.text, texts[-1])
        self.assertEqual(result.output_path, job.output_file.path)
        self.assertEqual(job.output_file.name, f'jobs/{job.pk}/output/translation.docx')
        self.assertEqual(
            [p.text for p in Document(result.output_path).paragraphs],
            [value for text in texts for value in ('TRADUIT PAR ZENIA', text)],
        )
        self.send.assert_called_once()
        email = self.send.call_args.args[0]
        self.assertEqual(email.to, [self.user.email])
        self.assertEqual(
            [p.text for p in Document(BytesIO(email.attachments[0][1])).paragraphs],
            [p.text for p in Document(result.output_path).paragraphs],
        )
        self.translate.assert_has_calls([call(text, job.source_language, is_pdf=True) for text in texts])
        self.assertEqual(self.translate.call_count, len(texts))

    def assert_rasterized_pages(self, job, numbers):
        workspace = workspace_for(job)
        self.assertEqual(self.convert.call_args_list, [
            call(str(workspace.source_dir / 'source.pdf'), dpi=300,
                 output_folder=str(workspace.working_dir),
                 first_page=number, last_page=number, paths_only=True)
            for number in numbers
        ])
        self.assertEqual(self.ocr.call_args_list, [
            call(workspace.working_dir / f'page_{number:04d}.ppm', job.source_language, is_pdf=True)
            for number in numbers
        ])

    def test_digital_pdf_keeps_native_text_without_ocr_or_rasterization(self):
        texts = ['Épilogue : où étions-nous ?', 'The story begins here.',
                 'Capitolo primo: perché?', '¿Dónde está el corazón?']
        job = self.native_job(texts)
        self.assert_document(job, process_document(job), texts)
        self.assert_rasterized_pages(job, [])
        self.assertEqual(len(list(workspace_for(job).working_dir.iterdir())), len(texts))

    def test_pdf_without_text_uses_ocr_for_each_page(self):
        job = self.native_job([None, None, None])
        self.ocr.side_effect = ['Première', 'Deuxième', 'Troisième']
        self.assert_document(job, process_document(job), ['Première', 'Deuxième', 'Troisième'])
        self.assert_rasterized_pages(job, [1, 2, 3])

    def test_scanned_pdf_with_only_an_embedded_image_uses_ocr(self):
        buffer = BytesIO()
        with Image.new('RGB', (300, 300), 'white') as image:
            ImageDraw.Draw(image).text((10, 20), 'Scanned chapter', fill='black')
            image.save(buffer, format='PDF')
        source = DocFile.objects.create(
            user=self.user, type_file='PDF', type_language='Anglais',
            button_televerse=SimpleUploadedFile('scan.pdf', buffer.getvalue(), content_type='application/pdf'),
        )
        job = TranslationJob.objects.create(user=self.user, source=source, source_language='Anglais')
        self.assert_document(job, process_document(job), ['Source'])
        self.assert_rasterized_pages(job, [1])

    def test_hybrid_pdf_rasterizes_only_missing_pages_and_preserves_order(self):
        job = self.native_job(['Page une', None, 'Page trois', None, 'Page cinq'])
        self.ocr.side_effect = ['Page deux', 'Page quatre']
        self.assert_document(job, process_document(job),
                             ['Page une', 'Page deux', 'Page trois', 'Page quatre', 'Page cinq'])
        self.assert_rasterized_pages(job, [2, 4])

    def test_nonempty_corrupt_native_text_triggers_ocr(self):
        job = self.native_job(['??? !!!', '(cid:123)', 'aaaaaaaa', '\x01\x02'])
        self.assert_document(job, process_document(job), ['Source'] * 4)
        self.assert_rasterized_pages(job, [1, 2, 3, 4])

    def test_short_native_pages_are_kept(self):
        texts = ['I', '7', 'Prologue', 'Capítulo 1']
        job = self.native_job(texts)
        self.assert_document(job, process_document(job), texts)
        self.assert_rasterized_pages(job, [])

    def test_japanese_and_chinese_extraction_is_preserved_verbatim(self):
        for language, text in [('Japonais', '第一章\n物語が始まる。'),
                               ('Chinois', '第一章\n故事開始了。')]:
            with self.subTest(language=language):
                self.translate.reset_mock()
                self.send.reset_mock()
                job = self.job(language=language)
                with patch('PyPDF2._page.PageObject.extract_text', return_value=text):
                    self.assert_document(job, process_document(job), [text, text])
                self.assert_rasterized_pages(job, [])

    def test_page_read_error_falls_back_only_for_that_page(self):
        job = self.native_job(['Une', None, 'Trois'])
        with patch('PyPDF2._page.PageObject.extract_text',
                   side_effect=['Une', PdfReadError('broken text stream'), 'Trois']):
            self.assert_document(job, process_document(job), ['Une', 'Source', 'Trois'])
        self.assert_rasterized_pages(job, [2])

    def test_unreadable_pdf_records_expected_error_and_preserves_cause(self):
        for stage in ['open', 'pages']:
            with self.subTest(stage=stage):
                job = self.job()
                error = PdfReadError('private PDF details')
                with patch('Tsukiyomi_app.services.extraction.PdfReader') as reader:
                    if stage == 'open':
                        reader.side_effect = error
                    else:
                        reader.return_value.pages.__iter__.side_effect = error
                    with self.assertRaises(DocumentProcessingError) as caught:
                        process_document(job)
                self.assertIs(caught.exception.__cause__, error)
                job.refresh_from_db()
                self.assertEqual(job.status, TranslationJob.Status.FAILED)
                self.assertEqual(job.error, 'PdfReadError')
                self.assertIsNotNone(job.finished_at)
                self.assertFalse(job.output_file)
                self.assert_rasterized_pages(job, [])
                self.translate.assert_not_called()
                self.send.assert_not_called()

    def test_unexpected_extraction_errors_are_recorded_and_reraised(self):
        for error in [ValueError('bug'), TypeError('bug'), KeyError('bug')]:
            with self.subTest(error=type(error).__name__):
                job = self.job()
                with patch('PyPDF2._page.PageObject.extract_text', side_effect=error):
                    with self.assertRaises(type(error)) as caught:
                        process_document(job)
                self.assertIs(caught.exception, error)
                job.refresh_from_db()
                self.assertEqual(job.status, TranslationJob.Status.FAILED)
                self.assertEqual(job.error, type(error).__name__)
                self.assertIsNotNone(job.finished_at)
                self.assertFalse(job.output_file)
                self.assert_rasterized_pages(job, [])
                self.send.assert_not_called()

    def test_ocr_failure_after_native_page_never_sends_partial_result(self):
        job = self.native_job(['Une', None, 'Trois'])
        error = TesseractError(1, 'private OCR details')
        self.ocr.side_effect = error
        with self.assertRaises(DocumentProcessingError) as caught:
            process_document(job)
        self.assertIs(caught.exception.__cause__, error)
        job.refresh_from_db()
        self.assertEqual(job.status, TranslationJob.Status.FAILED)
        self.assertEqual(job.error, 'TesseractError')
        self.assertIsNotNone(job.finished_at)
        self.assertFalse(job.output_file)
        self.assert_rasterized_pages(job, [2])
        self.translate.assert_called_once_with('Une', 'Anglais', is_pdf=True)
        self.assertEqual(list(workspace_for(job).output_dir.iterdir()), [])
        self.send.assert_not_called()

    def test_native_and_hybrid_jobs_remain_isolated(self):
        first = self.native_job(['Premier'])
        second = self.native_job(['Deuxième', None])
        process_document(first)
        snapshot = self.snapshot(workspace_for(first).root)
        process_document(second)
        self.assertEqual(self.snapshot(workspace_for(first).root), snapshot)
        self.assertNotEqual(first.output_file.name, second.output_file.name)
        self.assert_rasterized_pages(second, [2])
        with self.assertRaises(ValueError):
            process_document(TranslationJob.objects.get(pk=first.pk))

    def test_http_native_success_and_extraction_failure_preserve_responses(self):
        self.client.force_login(self.user)
        data = {'profileType': 'pdf_img', 'pi-type_file': 'PDF', 'pi-type_language': 'Anglais'}
        response = self.client.post(reverse('televerse_url'), {
            **data, 'pi-button_televerse': native_pdf_upload(['Native']),
        })
        self.assertContains(response, 'Native')
        self.assertTemplateUsed(response, 'tsukiyomi_app/succes_uploadfile.html')
        job = TranslationJob.objects.get()
        self.assertEqual(job.status, TranslationJob.Status.SUCCEEDED)
        self.assert_rasterized_pages(job, [])
        self.send.assert_called_once()
        self.send.reset_mock()
        with patch('Tsukiyomi_app.services.extraction.PdfReader', side_effect=[
            PdfReader(native_pdf_upload(['Native'])), PdfReadError('private'),
        ]):
            response = self.client.post(reverse('televerse_url'), {
                **data, 'pi-button_televerse': native_pdf_upload(['Native']),
            })
        self.assertContains(response, 'traitement', status_code=502)
        self.assertTemplateUsed(response, 'tsukiyomi_app/televerse_page.html')
        job = TranslationJob.objects.latest('created_at')
        self.assertEqual(job.status, TranslationJob.Status.FAILED)
        self.assertEqual(job.error, 'PdfReadError')
        self.assertFalse(job.output_file)
        self.assertEqual(DocFile.objects.count(), 2)
        self.send.assert_not_called()
