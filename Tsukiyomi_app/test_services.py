import os
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase
from docx import Document
from pytesseract.pytesseract import TesseractError

from .models import DocFile, TranslationJob
from .services.errors import DocumentProcessingError
from .services.processor import process_document
from .services.workspace import workspace_for
from .tests import image_upload, pdf_upload


class DocumentProcessorTests(TestCase):
    """Appeler le pipeline sans HTTP, avec jobs persistants et DOCX réels."""

    def setUp(self):
        temporary = TemporaryDirectory(prefix='tsukiyomi-services-')
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        previous = Path.cwd()
        os.chdir(self.root)
        self.addCleanup(os.chdir, previous)
        config = self.settings(MEDIA_ROOT=self.root)
        config.enable()
        self.addCleanup(config.disable)
        self.user = get_user_model().objects.create_user(username='reader', email='reader@example.test')
        self.start_patch('socket.socket.connect', side_effect=AssertionError('Réseau interdit'))
        self.start_patch('subprocess.Popen', side_effect=AssertionError('Processus externe interdit'))
        self.convert = self.start_patch('Tsukiyomi_app.services.extraction.convert_from_path', return_value=['page'])
        self.ocr = self.start_patch('Tsukiyomi_app.services.ocr.pytesseract.image_to_string', return_value='Source')
        self.translator = self.start_patch('Tsukiyomi_app.services.translation.GoogleTranslator')
        self.translator.return_value.translate.return_value = 'Traduction'
        self.send = self.start_patch('Tsukiyomi_app.services.notification.EmailMessage.send',
                                     autospec=True, return_value=1)

    def start_patch(self, target, **kwargs):
        patcher = patch(target, **kwargs)
        result = patcher.start()
        self.addCleanup(patcher.stop)
        return result

    def job(self, kind='PDF', language='Anglais', pages=1):
        upload = pdf_upload(pages) if kind == 'PDF' else image_upload()
        document = DocFile.objects.create(
            user=self.user, type_file=kind, type_language=language, button_televerse=upload,
        )
        return TranslationJob.objects.create(user=self.user, source=document, source_language=language)

    def test_all_languages_preserve_external_options_and_docx_without_http(self):
        for kind in ['PDF', 'Image']:
            for language, ocr_code, translation_code in [
                ('Français', 'fra', 'fr'), ('Anglais', 'eng', 'en'),
                ('Italien', 'ita', 'it'), ('Espagnol', 'spa', 'es'),
                ('Japonais', 'jpn', 'ja'), ('Chinois', 'chi_tra', 'zh-TW'),
            ]:
                with self.subTest(kind=kind, language=language):
                    for mock in [self.convert, self.ocr, self.translator, self.send]:
                        mock.reset_mock()
                    job = self.job(kind, language)
                    result = process_document(job)
                    self.assertEqual(result.text, 'Traduction')
                    self.assertTrue(Path(result.output_path).is_file())
                    expected_ocr = {'lang': ocr_code}
                    if kind == 'PDF' or language == 'Français':
                        expected_ocr['config'] = '--oem 3 --psm 6'
                    self.assertEqual(self.ocr.call_args.kwargs, expected_ocr)
                    self.translator.assert_called_once_with(source=translation_code, target='fr')
                    options = {} if kind == 'PDF' and language == 'Français' else {'timeout': 10}
                    self.translator.return_value.translate.assert_called_once_with('Source', **options)
                    if kind == 'PDF':
                        workspace = workspace_for(job)
                        self.convert.assert_called_once_with(str(workspace.source_dir / 'source.pdf'), dpi=300,
                                                             output_folder=str(workspace.working_dir),
                                                             first_page=1, last_page=1, paths_only=True)
                    else:
                        self.convert.assert_not_called()
                        self.assertIsNone(self.ocr.call_args.args[0].fp)
                    self.send.assert_called_once()
                    email = self.send.call_args.args[0]
                    self.assertEqual(email.to, [self.user.email])
                    self.assertEqual(email.from_email, settings.DEFAULT_FROM_EMAIL)
                    self.assertEqual(email.subject, f'Votre fichier de langue source {language} traduit en Français')
                    self.assertEqual(email.body, 'Veuillez trouver ci-joint votre fichier traduit par TSUKIYOMI')
                    self.assertFalse(self.send.call_args.kwargs['fail_silently'])
                    self.assertEqual(len(email.attachments), 1)
                    self.assertEqual(email.attachments[0][0], Path(result.output_path).name)
                    doc = Document(BytesIO(email.attachments[0][1]))
                    self.assertEqual([p.text for p in doc.paragraphs], ['TRADUIT PAR ZENIA', 'Traduction'])

    def test_pdf_result_contains_last_translation_and_all_pages_in_order(self):
        self.convert.side_effect = [['page-0'], ['page-1'], ['page-2']]
        self.translator.return_value.translate.side_effect = ['PREMIERE', 'DEUXIEME', 'DERNIERE']
        result = process_document(self.job(pages=3))
        self.assertEqual(result.text, 'DERNIERE')
        doc = Document(result.output_path)
        self.assertEqual([p.text for p in doc.paragraphs if p.text != 'TRADUIT PAR ZENIA'],
                         ['PREMIERE', 'DEUXIEME', 'DERNIERE'])
        self.send.assert_called_once()

    def test_empty_later_page_returns_no_final_artifact_or_email(self):
        for empty in ['', None]:
            with self.subTest(empty=empty):
                self.ocr.reset_mock()
                self.translator.reset_mock()
                self.ocr.side_effect = ['Source', empty, 'Ne doit pas être lu']
                result = process_document(self.job(pages=3))
                self.assertIsNone(result.text)
                self.assertIsNone(result.output_path)
                self.assertEqual(self.ocr.call_count, 2)
                self.translator.return_value.translate.assert_called_once()
                self.send.assert_not_called()
                self.assertFalse(list(self.root.rglob('output/*.docx')))

    def test_later_ocr_failure_preserves_cause_and_stops_before_merge(self):
        error = TesseractError(1, 'failure')
        self.ocr.side_effect = ['Source', error, 'Ne doit pas être lu']
        with self.assertRaises(DocumentProcessingError) as caught:
            process_document(self.job(pages=3))
        self.assertIs(caught.exception.__cause__, error)
        self.assertEqual(self.ocr.call_count, 2)
        self.translator.return_value.translate.assert_called_once()
        self.send.assert_not_called()
        self.assertFalse(list(self.root.rglob('output/*.docx')))

    def test_image_is_closed_on_expected_and_unexpected_translation_errors(self):
        for error in [OSError('offline'), ValueError('programming error')]:
            with self.subTest(error=type(error).__name__):
                self.translator.return_value.translate.side_effect = error
                expected = DocumentProcessingError if isinstance(error, OSError) else ValueError
                with self.assertRaises(expected) as caught:
                    process_document(self.job('Image'))
                if expected is DocumentProcessingError:
                    self.assertIs(caught.exception.__cause__, error)
                else:
                    self.assertIs(caught.exception, error)
                self.assertIsNone(self.ocr.call_args.args[0].fp)
                self.send.assert_not_called()
                self.assertFalse(list(self.root.rglob('*.docx')))

    def test_rendering_failure_stops_before_notification(self):
        error = OSError('disk full')
        for kind in ['PDF', 'Image']:
            with self.subTest(kind=kind), patch('docx.document.Document.save', side_effect=error):
                with self.assertRaises(DocumentProcessingError) as caught:
                    process_document(self.job(kind))
                self.assertIs(caught.exception.__cause__, error)
                self.send.assert_not_called()
