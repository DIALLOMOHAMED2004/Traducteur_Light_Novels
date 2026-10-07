from dataclasses import replace
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, call, patch

from deep_translator.exceptions import RequestError, TooManyRequests, TranslationNotFound
from django.test import SimpleTestCase, TestCase
from docx import Document
from requests.exceptions import Timeout

from .models import DocFile, TranslationJob
from .services.errors import DocumentProcessingError, document_errors
from .services.languages import LANGUAGES
from .services.processor import process_document
from .services.rendering import write_docx
from .services.segmentation import segment_text
from .services.translation import translate_segments, translate_text
from .services.workspace import workspace_for
from .test_extraction import native_pdf_upload
from .test_jobs import JobTestSetup


class SegmentationTests(SimpleTestCase):
    def test_paragraphs_keep_internal_lines_and_deterministic_identity(self):
        text = 'Premier paragraphe\nsur deux lignes.\n\nDeuxième.\n\nTroisième.'
        segments = segment_text(text, page_number=2)
        self.assertEqual(segments, segment_text(text, page_number=2))
        self.assertEqual([s.source_text for s in segments], [
            'Premier paragraphe\nsur deux lignes.', 'Deuxième.', 'Troisième.',
        ])
        self.assertEqual([s.id for s in segments], ['p0002-s0001', 'p0002-s0002', 'p0002-s0003'])
        self.assertEqual([s.order for s in segments], [1, 2, 3])
        self.assertTrue(all(s.page_number == 2 and s.kind == 'paragraph' for s in segments))
        self.assertTrue(all(s.translated_text is None for s in segments))
        self.assertEqual(segment_text(text)[0].id, 'p0001-s0001')

    def test_blank_lines_do_not_create_empty_segments(self):
        self.assertEqual(segment_text(' \n\r\n\t'), [])
        self.assertEqual(segment_text(''), [])
        segments = segment_text('\nPremier\r\n \t\r\n\r\nSecond\n\n')
        self.assertEqual([s.source_text for s in segments], ['Premier', 'Second'])
        self.assertEqual([s.order for s in segments], [1, 2])

    def test_explicit_structure_is_recognized_without_blank_lines(self):
        text = '# Chapitre 1\nRécit.\n— Bonjour !\n***\nSuite.'
        segments = segment_text(text)
        self.assertEqual([s.kind for s in segments], [
            'heading', 'paragraph', 'dialogue', 'separator', 'paragraph',
        ])
        self.assertEqual([s.source_text for s in segments], text.splitlines())

    def test_only_unambiguous_separator_lines_are_classified(self):
        for text in ['***', '* * *', '---', '___', '===', '※※※', '———']:
            with self.subTest(text=text):
                self.assertEqual(segment_text(text)[0].kind, 'separator')
        for text in ['...', '- élément', '* mot *', 'a --- b', '--- fin']:
            with self.subTest(text=text):
                self.assertEqual(segment_text(text)[0].kind, 'paragraph')

    def test_explicit_dialogue_lines_keep_their_markers(self):
        for text in ['— Bonjour.', '– Hello.', '« Bonjour »', '“Hello”', '「こんにちは」', '『你好』']:
            with self.subTest(text=text):
                segment = segment_text(text)[0]
                self.assertEqual(segment.kind, 'dialogue')
                self.assertEqual(segment.source_text, text)

    def test_uncertain_titles_and_dialogues_remain_paragraphs(self):
        for text in ['Bonjour.', 'CHAPITRE', 'Première ligne\nSuite', '「Citation ouverte',
                     'Il dit « Bonjour ».', '第一章\n物語が始まる。', '第一章\n故事開始了。']:
            with self.subTest(text=text):
                segments = segment_text(text)
                self.assertEqual(len(segments), 1)
                self.assertEqual(segments[0].kind, 'paragraph')
                self.assertEqual(segments[0].source_text, text)

    def test_japanese_and_chinese_do_not_require_word_spaces(self):
        for text, expected in [
            ('物語が始まる。\n続き。\n\n「こんにちは」\n\n終わり。',
             ['物語が始まる。\n続き。', '「こんにちは」', '終わり。']),
            ('故事開始了。\n繼續。\n\n『你好』\n\n結束。',
             ['故事開始了。\n繼續。', '『你好』', '結束。']),
        ]:
            with self.subTest(text=text):
                segments = segment_text(text)
                self.assertEqual([s.source_text for s in segments], expected)
                self.assertEqual([s.kind for s in segments], ['paragraph', 'dialogue', 'paragraph'])

    def test_no_significant_content_or_internal_spacing_is_lost(self):
        text = '\n  # Titre\r\n\r\n  E\u0301té\u00a0:  世界。\r\nSuite\tfin.\n— Oui !\n* * *\n'
        segments = segment_text(text)
        self.assertEqual(segments[1].source_text, '  E\u0301té\u00a0:  世界。\nSuite\tfin.')
        self.assertEqual(
            ''.join(''.join(s.source_text.split()) for s in segments),
            ''.join(text.split()),
        )


class SegmentTranslationTests(SimpleTestCase):
    def test_callable_provider_preserves_metadata_source_and_all_languages(self):
        source = segment_text('# Title\nFirst.\n\nSecond.\n***\n— Hello.')
        for language in LANGUAGES:
            for is_pdf in [False, True]:
                with self.subTest(language=language, is_pdf=is_pdf):
                    provider = Mock(side_effect=lambda text, *args, **kwargs: 'FR ' + text)
                    translated = translate_segments(source, language, provider=provider, is_pdf=is_pdf)
                    self.assertEqual(provider.call_args_list, [
                        call(s.source_text, language, is_pdf=is_pdf)
                        for s in source if s.kind != 'separator'
                    ])
                    self.assertEqual([replace(s, translated_text=None) for s in translated], source)
                    self.assertTrue(all(s.translated_text is None for s in source))
                    self.assertEqual([s.translated_text for s in translated], [
                        'FR # Title', 'FR First.', 'FR Second.', '***', 'FR — Hello.',
                    ])

    def test_separator_only_never_calls_provider(self):
        provider = Mock(side_effect=AssertionError('Aucun appel attendu'))
        translated = translate_segments(segment_text('***'), 'Anglais', provider=provider)
        self.assertEqual(translated[0].translated_text, '***')
        provider.assert_not_called()

    def test_google_errors_keep_the_existing_cause_and_stop_at_failing_segment(self):
        for error in [RequestError(), TooManyRequests(), TranslationNotFound('source'), Timeout()]:
            with self.subTest(error=type(error).__name__), patch(
                'Tsukiyomi_app.services.translation.GoogleTranslator',
            ) as google:
                google.return_value.translate.side_effect = ['Un', error, 'Trois']
                with self.assertRaises(DocumentProcessingError) as caught, document_errors():
                    translate_segments(segment_text('One\n\nTwo\n\nThree'), 'Anglais', provider=translate_text)
                self.assertIs(caught.exception.__cause__, error)
                self.assertEqual(google.return_value.translate.call_count, 2)

    def test_unexpected_provider_error_is_not_hidden(self):
        error = ValueError('bug')
        with self.assertRaises(ValueError) as caught, document_errors():
            translate_segments(segment_text('Texte'), 'Anglais', provider=Mock(side_effect=error))
        self.assertIs(caught.exception, error)


class SegmentRenderingTests(SimpleTestCase):
    def test_real_docx_contains_translations_types_and_separator_in_supplied_order(self):
        source = segment_text('# Title\nFirst\n\nSecond\n— Hello\n***\nLast')
        segments = [replace(s, translated_text=f'Traduction {s.order}') for s in source]
        with TemporaryDirectory(prefix='tsukiyomi-segments-') as temporary:
            path = Path(temporary) / 'translation.docx'
            write_docx(segments, path)
            doc = Document(path)
        self.assertEqual([p.text for p in doc.paragraphs], [
            'TRADUIT PAR ZENIA', 'Traduction 1', 'Traduction 2', 'Traduction 3',
            'Traduction 4', '***', 'Traduction 6',
        ])
        self.assertEqual([p.style.name for p in doc.paragraphs[1:]], [
            'Heading 1', 'Normal', 'Normal', 'Quote', 'Normal', 'Normal',
        ])

    def test_missing_translation_does_not_silently_render_source(self):
        with TemporaryDirectory(prefix='tsukiyomi-segments-') as temporary:
            path = Path(temporary) / 'translation.docx'
            with self.assertRaisesRegex(ValueError, 'sans traduction'):
                write_docx(segment_text('Source'), path)
            self.assertFalse(path.exists())


class SegmentedPipelineTests(JobTestSetup, TestCase):
    """Provider injecté, PDF/DOCX réels, OCR et SMTP simulés, réseau interdit."""

    def test_digital_scanned_hybrid_and_image_render_segments_in_page_order(self):
        first = '# Title\nFirst\n\nSecond\n— Hello\n***'
        last = 'Last\n\nEnd'
        for scenario in ['digital', 'scanned', 'hybrid', 'image']:
            with self.subTest(scenario=scenario):
                self.convert.reset_mock()
                self.ocr.reset_mock()
                self.send.reset_mock()
                if scenario == 'image':
                    job = self.job('Image')
                    self.ocr.side_effect = [first]
                    pages = [first]
                else:
                    native = [first, last] if scenario == 'digital' else (
                        [None, None] if scenario == 'scanned' else [first, None]
                    )
                    source = DocFile.objects.create(
                        user=self.user, type_file='PDF', type_language='Anglais',
                        button_televerse=native_pdf_upload(native),
                    )
                    job = TranslationJob.objects.create(user=self.user, source=source, source_language='Anglais')
                    self.ocr.side_effect = [first, last] if scenario == 'scanned' else [last]
                    pages = [first, last]
                provider = Mock(side_effect=lambda text, *args, **kwargs: 'FR ' + text)
                with patch('Tsukiyomi_app.services.processor.write_docx', wraps=write_docx) as render, patch(
                    'Tsukiyomi_app.services.translation.GoogleTranslator',
                    side_effect=AssertionError('Google ne doit pas être instancié'),
                ):
                    result = process_document(job, provider=provider)
                expected_segments = [
                    replace(s, translated_text=s.source_text if s.kind == 'separator' else 'FR ' + s.source_text)
                    for number, text in enumerate(pages, start=1)
                    for s in segment_text(text, page_number=number)
                ]
                self.assertEqual([s for args in render.call_args_list for s in args.args[0]], expected_segments)
                self.assertEqual(provider.call_args_list, [
                    call(s.source_text, 'Anglais', is_pdf=scenario != 'image')
                    for s in expected_segments if s.kind != 'separator'
                ])
                expected_texts = [s.translated_text for s in expected_segments]
                self.assertEqual([p.text for p in Document(result.output_path).paragraphs
                                  if p.text != 'TRADUIT PAR ZENIA'], expected_texts)
                self.send.assert_called_once()
                attachment = Document(BytesIO(self.send.call_args.args[0].attachments[0][1]))
                self.assertEqual([p.text for p in attachment.paragraphs],
                                 [p.text for p in Document(result.output_path).paragraphs])
                self.assertEqual(result.text, '\n\n'.join(s.translated_text for s in expected_segments
                                                         if s.page_number == len(pages)))
                self.assertEqual(self.convert.call_count, {'digital': 0, 'scanned': 2, 'hybrid': 1, 'image': 0}[scenario])
                if scenario == 'hybrid':
                    self.assertEqual(self.convert.call_args.kwargs['first_page'], 2)
                job.refresh_from_db()
                self.assertEqual(job.status, TranslationJob.Status.SUCCEEDED)
                self.assertEqual(result.output_path, job.output_file.path)
                self.translate.assert_not_called()

    def test_later_segment_failure_never_sends_partial_result(self):
        for kind in ['PDF', 'Image']:
            with self.subTest(kind=kind):
                job = self.job(kind)
                self.ocr.side_effect = ['Previous', 'One\n\nTwo\n\nThree'] if kind == 'PDF' else ['One\n\nTwo\n\nThree']
                error = RequestError()
                replies = ['Previous', 'Un', error] if kind == 'PDF' else ['Un', error]
                provider = Mock(side_effect=replies)
                with self.assertRaises(DocumentProcessingError) as caught:
                    process_document(job, provider=provider)
                self.assertIs(caught.exception.__cause__, error)
                self.assertEqual(provider.call_count, len(replies))
                job.refresh_from_db()
                self.assertEqual(job.status, TranslationJob.Status.FAILED)
                self.assertEqual(job.error, 'RequestError')
                self.assertFalse(job.output_file)
                self.assertEqual(list(workspace_for(job).output_dir.iterdir()), [])
                self.send.assert_not_called()

    def test_whitespace_only_page_has_no_output_or_notification(self):
        for kind in ['PDF', 'Image']:
            with self.subTest(kind=kind):
                job = self.job(kind)
                self.ocr.side_effect = ['Previous', ' \n\t'] if kind == 'PDF' else [' \n\t']
                result = process_document(job)
                self.assertIsNone(result.output_path)
                self.assertIsNone(result.text)
                self.assertEqual(list(workspace_for(job).output_dir.iterdir()), [])
                self.send.assert_not_called()
