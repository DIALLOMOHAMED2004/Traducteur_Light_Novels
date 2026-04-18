from PIL import Image
import pytesseract
from deep_translator import GoogleTranslator
from docx import Document
from docx2pdf import convert

#code pour extraire le text de l'image
img = Image.open('Medias_test\pack actualisé.png')
text = pytesseract.image_to_string(img, lang='fra')

print(text)

#code pour traduire en francais le text extrait
# translated_fr = GoogleTranslator(source='en', target='fr').translate(text)

# #print(translated_fr)

# #code pour recuperer le text traduit et l'inserer dans un fichier docx
# doc = Document()
# doc.add_heading('VOICI LA PREMIERE TRADUCTION EN FRANCAIS')
# doc.add_paragraph(translated_fr)
# doc.add_picture('Medias_test\page 1.PNG')
# doc.save('trad_page1.docx')


# #code qui permet de convertir le fichier docx en pdf
# convert("trad_page1.docx", "trad_page1.pdf")


