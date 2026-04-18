from PIL import Image
import pytesseract
from deep_translator import GoogleTranslator
from docx import Document
from docx2pdf import convert
import cv2

# Ouvre la caméra (0 = caméra par défaut)
cap = cv2.VideoCapture(0)



while True:
    ret, frame = cap.read()
    cv2.imshow("TSUKIYOMI - Caméra", frame)
    # gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    # blur = cv2.GaussianBlur(gray, (5, 5), 0)
    # _, thresh = cv2.threshold(blur, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)


    if cv2.waitKey(1) & 0xFF == ord('s'):  # appuie sur 's' pour capturer
        cv2.imwrite("IMAGE_CAPTURE\\capture_clean.png", frame)
        break

cap.release()
cv2.destroyAllWindows()



#code pour extraire le text de l'image
img = Image.open("IMAGE_CAPTURE\\capture_clean.png")
text = pytesseract.image_to_string(img, lang='fra')
print(text)

# #code pour traduire en francais le text extrait
# translated_fr = GoogleTranslator(source='en', target='fr').translate(text)

# #print(translated_fr)

# #code pour recuperer le text traduit et l'inserer dans un fichier docx
# doc = Document()
# doc.add_heading('VOICI LA PREMIERE TRADUCTION EN FRANCAIS')
# doc.add_paragraph(translated_fr)
# #doc.add_picture('Medias_test\\test2.PNG')
# doc.save('trad_page1.docx')


# #code qui permet de convertir le fichier docx en pdf
# convert("trad_page1.docx", "trad_page1.pdf")



















